"""
Script to generate the complete RegolithSense Solution Concept Note as a formatted Word Document (.docx).
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color.replace('#', ''))
    tcPr.append(shd)

def create_concept_note_docx(output_path):
    doc = docx.Document()
    
    # Page Margins (0.75 in)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 1. Header Banner
    prog_p = doc.add_paragraph()
    prog_run = prog_p.add_run("PAN-AFRICAN MINETECH EXPLORATION TECHNOLOGY PRE-INCUBATION PROGRAMME\nNIGERIA TRACK — CRITICAL MINERALS CHALLENGE")
    prog_run.font.name = "Arial"
    prog_run.font.size = Pt(8.5)
    prog_run.font.bold = True
    prog_run.font.color.rgb = RGBColor(100, 116, 139) # Slate 500
    prog_p.paragraph_format.space_after = Pt(4)

    # Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("SOLUTION CONCEPT NOTE")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    title_p.paragraph_format.space_after = Pt(2)

    # Subtitle
    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("RegolithSense Enterprise: A Cloud-Native Multi-Physics Inversion and Spatial AI Prospectivity Platform for Deeply Concealed Critical Mineral Systems Beneath Tropical Weathering Cover")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(11)
    sub_run.font.bold = True
    sub_run.font.color.rgb = RGBColor(2, 132, 199) # Sky Blue 600
    sub_p.paragraph_format.space_after = Pt(14)

    # Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_info = [
        ("Target Challenge Brief:", "Concealed Subsurface Geology Beneath Deep Regolith / Overburden Cover (15m–80m)"),
        ("Organizers & Partners:", "timbuktoo MineTech Hub & United Nations Development Programme (UNDP) in partnership with UniPod Nigeria"),
        ("Lead Institutional Affiliation:", "Nasarawa State University Keffi (NSUK) Mine-Tech UniPod / Federal University of Lafia"),
        ("Primary Exploration Targets:", "Lithium-Cesium-Tantalum (LCT) Pegmatites, Columbite-Tantalite & Rare Earth Elements (REEs)")
    ]
    for row_idx, (k, v) in enumerate(meta_info):
        cell_k = meta_table.cell(row_idx, 0)
        cell_v = meta_table.cell(row_idx, 1)
        cell_k.text = k
        cell_v.text = v
        
        pk = cell_k.paragraphs[0]
        pk.runs[0].font.name = "Arial"
        pk.runs[0].font.size = Pt(9)
        pk.runs[0].font.bold = True
        set_cell_background(cell_k, "#F1F5F9")
        
        pv = cell_v.paragraphs[0]
        pv.runs[0].font.name = "Arial"
        pv.runs[0].font.size = Pt(9)
        set_cell_background(cell_v, "#FAFAFA")
    
    doc.add_paragraph() # Spacer

    def add_section_heading(text):
        h = doc.add_heading(text, level=1)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        for r in h.runs:
            r.font.name = "Arial"
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)

    def add_body_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Arial"
            r_bold.font.size = Pt(9.5)
            r_bold.font.bold = True
            r_bold.font.color.rgb = RGBColor(15, 23, 42)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(51, 65, 85)
        return p

    # 1. Executive Summary & Problem Formulation
    add_section_heading("1. Executive Summary & Problem Formulation")
    add_body_p(
        "A fundamental paradox defines mineral exploration across tropical Africa: while the continent hosts globally significant endowments of critical transition minerals—including lithium-cesium-tantalum (LCT) pegmatites, rare earth elements (REEs), columbite-tantalite, nickel, and copper—the overwhelming majority of these prospective bedrock formations remain entirely concealed beneath deep tropical regolith cover, thick lateritic duricrusts, saprolite, and dense vegetation blankets. Over tens of millions of years of intense sub-aerial chemical weathering, overburden profiles ranging from 15 to over 80 meters in thickness have developed across the Pan-African Basement Complex and surrounding greenstone-schist belts."
    )
    add_body_p(
        "This deep weathered mantle invalidates standard mineral exploration workflows:", bold_prefix="The Technical Bottleneck: "
    )
    add_body_p(
        "Conventional optical and multispectral remote sensing techniques (such as Landsat, Sentinel-2, and ASTER band ratios) are physically limited to sensing only the top few micrometers of the Earth's surface, reflecting barren ferricrete, agricultural soil, or dense vegetative canopy rather than mineralized bedrock.",
        bold_prefix="• Optical Satellites Fail: "
    )
    add_body_p(
        "Surface soil geochemical signatures are severely leached, mechanically transported, or displaced downslope by tropical drainage, generating costly false positives and masking blind ore bodies.",
        bold_prefix="• Geochemical Dispersion: "
    )
    add_body_p(
        "Junior and mid-tier exploration companies operating in Nigeria face astronomical exploration risks, routinely expending up to 50% of their early-stage capital on arbitrary, unconstrained exploratory grid drilling ($150–$250/meter) or prematurely abandoning highly prospective greenfield tenements.",
        bold_prefix="• Capital Destruction: "
    )
    add_body_p(
        "There is an urgent, unmet market demand for an integrated computational exploration platform that can mathematically penetrate deep regolith cover, filter transported surficial noise, and reconstruct concealed bedrock geology prior to committing drilling capital."
    )

    # 2. Proposed Technological Innovation
    add_section_heading("2. Proposed Technological Innovation: RegolithSense Enterprise")
    add_body_p(
        "RegolithSense Enterprise is an innovative, cloud-enabled multi-physics geophysical inversion and spatial machine learning platform engineered specifically to 'see through' thick tropical regolith and delineate concealed critical mineral systems. By combining airborne potential field geophysics, satellite-derived gravity models, airborne gamma-ray spectrometry, and evidence-weighted machine learning, RegolithSense generates high-resolution, drill-ready prospectivity indices beneath tens of meters of barren overburden."
    )
    add_body_p("The innovation integrates five proprietary computational layers:")

    add_body_p(
        "High-resolution aeromagnetic grids are subjected to advanced frequency-domain filtering—including First Vertical Derivatives (1VD via 2D FFT), Total Horizontal Derivatives (THDR), Analytic Signal (AS), and the Tilt Angle Derivative (TDR; Miller & Singh, 1994). Because magnetic fields penetrate non-magnetic regolith without attenuation, these operators map deep ductile shear zones, brittle faults, lithological sutures, and buried pegmatite dyke swarms obscured beneath laterite.",
        bold_prefix="• Sub-Regolith Magnetic Wavefield Inversion: "
    )
    add_body_p(
        "Ingests high-resolution satellite-derived gravity data (GGMplus ~200m spatial resolution and WGM2012) combined with regional airborne gravity archives. Felsic peraluminous leucogranites (the parent magmas of lithium pegmatites) possess lower bulk density (ρ ≈ 2.58–2.64 g/cm³) compared to dense host migmatites and amphibolites (ρ ≈ 2.78–3.05 g/cm³). Negative Bouguer gravity anomaly lows and the First Vertical Gravity Gradient (Gzz in Eotvos) pinpoint these fertile intrusive roots beneath regolith cover with zero ground acquisition costs.",
        bold_prefix="• Satellite-Derived Bouguer Gravity & Density Inversion: "
    )
    add_body_p(
        "Evaluates Total Count (TC) and individual radioelements (K %, eTh ppm, eU ppm). Computes diagnostic exploration ratios (K/eTh, eU/eTh, eU/K) and the Efimov / Saunders F-Parameter Alteration Index (F = K * eU / eTh). This isolates hydrothermal potassic metasomatism (microcline, muscovite, lepidolite) typical of LCT pegmatite swarms, successfully distinguishing fresh in-situ bedrock weathering from barren transported laterite gravels.",
        bold_prefix="• Gamma-Ray Spectrometry & Alteration Suite: "
    )
    add_body_p(
        "A spatial artificial intelligence engine supporting Random Forest, Gradient Boosting, and Weights of Evidence (WofE) ensemble classifiers. Stacks structural proximity, magnetic contrast, gravity density lows, and radiometric indices into a unified predictive model that quantifies exploration uncertainty, validates with ROC-AUC scores (>0.92), and ranks exploration drill targets into Tier-1, Tier-2, and Tier-3 categories.",
        bold_prefix="• Machine Learning Mineral Prospectivity Mapping (MPM): "
    )
    add_body_p(
        "Renders 3D digital elevation models (DEM) with draped prospectivity models, East-West cutaway volume slicing, and subterranean drill string trajectories. Features a drag-and-drop survey importer (CSV/XYZ) for exploration teams and a one-click automated Word (.docx) and Excel (.xlsx) technical report generator.",
        bold_prefix="• 3D Subsurface Profiling & Enterprise Reporting: "
    )

    # 3. Technical Workflow & Data Processing Pipeline
    add_section_heading("3. Technical Workflow & Processing Pipeline")
    add_body_p("The RegolithSense technical pipeline executes across four standardized operational phases:")
    add_body_p(
        "Ingestion, reprojection, and Delaunay triangulation interpolation of regional airborne geophysics (100–500m line spacing), GGMplus satellite gravity, and digital elevation models into a standardized raster database.",
        bold_prefix="• Phase 1 — Multi-Source Geospatial Harmonization: "
    )
    add_body_p(
        "Automated classification of regolith regimes into: (a) Relict lateritic plateaus; (b) Erosional stripped saprolite zones; and (c) Depositional transported alluvial channels, preventing false geochemical anomalies.",
        bold_prefix="• Phase 2 — Regolith-Landform Architecture Mapping (RLM): "
    )
    add_body_p(
        "Automated edge detection algorithms extract concealed deep-seated fault architectures, shear zones, and granite margins from first vertical derivative, tilt angle derivative, and horizontal gravity gradient grids.",
        bold_prefix="• Phase 3 — Deep Structural & Density Extraction: "
    )
    add_body_p(
        "Supervised machine learning fuses evidential layers into continuous mineral prospectivity heatmaps (0% to 100%). Prospective zones are ranked from Tier-1 (immediate ground RC drilling) to Tier-3, shrinking concession search spaces by up to 80% and exporting JORC/NI 43-101 drill spotting packages.",
        bold_prefix="• Phase 4 — Multi-Criteria Stacking & Target Prioritization: "
    )

    # 4. Commercial Viability, Target Market, and Business Model
    add_section_heading("4. Commercial Viability & Business Model")
    add_body_p(
        "Greenfield mineral exploration across Africa suffers from prolonged discovery cycles, frequently exceeding 7 to 12 years from concession acquisition to maiden discovery. RegolithSense drastically shortens this timeline, offering compelling commercial value across key mining ecosystem participants:"
    )
    add_body_p(
        "Reduces concession holding costs and saves up to 50% on early-stage exploration budgets by eliminating blind grid drilling across barren transported cover, directing capital strictly to high-probability structural targets.",
        bold_prefix="• Junior Exploration & Mid-Tier Mining Companies: "
    )
    add_body_p(
        "Empowers government agencies (such as the Nigerian Geological Survey Agency - NGSA and the Solid Minerals Development Fund - SMDF) to systematically reprocess legacy national airborne geophysical data, unlocking massive latent asset value and attracting foreign direct investment (FDI).",
        bold_prefix="• National Geological Surveys & Mining Ministries: "
    )
    add_body_p(
        "Provides rapid, independent technical due diligence and asset valuation on exploration licenses undergoing acquisition or farm-in negotiations.",
        bold_prefix="• Mineral Investment Funds & Private Equity: "
    )
    add_body_p(
        "Commercial revenue is captured through: (1) Enterprise SaaS Platform Subscriptions ($1,500–$5,000/month/concession); (2) Multi-Client Regional Prospectivity Reports covering prolific belts (Central Nigerian Lithium-Pegmatite Belt and the Schist Belts); and (3) Tailored Turnkey Exploration Advisory Contracts."
    )

    # 5. Environmental, Social, and Governance (ESG) Impact
    add_section_heading("5. Environmental, Social, and Governance (ESG) Impact")
    add_body_p(
        "Traditional greenfield exploration through thick regolith involves extensive bulldozer trenching, road clearance, and unguided exploratory grid drilling. By narrowing drill targets from hundreds of square kilometers down to precise structural corridors, RegolithSense preserves delicate tropical ecosystems, prevents deforestation, and protects topsoil integrity (SDG 15: Life on Land).",
        bold_prefix="• Drastic Reduction in Environmental Footprint: "
    )
    add_body_p(
        "Accelerates the discovery pipeline for critical green energy minerals (lithium, tantalum, rare earths, nickel, copper) necessary for energy storage and electric mobility, supporting global decarbonization commitments (SDG 7: Affordable and Clean Energy; SDG 13: Climate Action).",
        bold_prefix="• Responsible Sourcing for the Global Energy Transition: "
    )
    add_body_p(
        "Builds African data-processing sovereignty, enabling domestic geoscientists to extract actionable commercial intelligence from national data assets rather than exporting raw data for expensive overseas consulting analysis (SDG 9: Industry, Innovation, and Infrastructure).",
        bold_prefix="• Indigenous Technological Sovereignty: "
    )

    # 6. Current Stage & 6-Month Incubation Roadmap
    add_section_heading("6. Current Stage of Development & 6-Month Roadmap")
    add_body_p(
        "RegolithSense Enterprise is currently at the Functional Applied Prototype Stage (TRL 4–5), operating on local/cloud instances with full 3D subsurface profiling, satellite gravity inversion, and automated Word/Excel reporting modules."
    )

    # Roadmap Table
    road_table = doc.add_table(rows=4, cols=2)
    road_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    phases = [
        ("Months 1 – 2:\nData Harmonization & Importer", "Complete automated CSV/XYZ/GeoTIFF raster ingestion engine; integrate GGMplus 200m satellite gravity and NGSA airborne magnetic leveling across the Central Nasarawa corridor."),
        ("Months 3 – 4:\nAlgorithm Training & UniPod Ground-Truthing", "Train supervised Random Forest and WofE prospectivity classifiers against verified mineral deposits; conduct ground validation using field pXRF and gamma-ray scintillometers across pilot regolith test sites in partnership with the NSUK Mine-Tech UniPod."),
        ("Months 5 – 6:\nCommercial Beta & Investor Showcase", "Deploy commercial beta with two junior mining operators in North-Central Nigeria; deliver a pilot regional prospectivity atlas and pitch the commercial venture to international mining investors at the timbuktoo MineTech Demo Day.")
    ]
    for r_idx, (phase_title, phase_desc) in enumerate(phases):
        c0 = road_table.cell(r_idx + 1, 0)
        c1 = road_table.cell(r_idx + 1, 1)
        c0.text = phase_title
        c1.text = phase_desc
        set_cell_background(c0, "#F8FAFC")
        set_cell_background(c1, "#FFFFFF")
        p0 = c0.paragraphs[0]
        p0.runs[0].font.name = "Arial"
        p0.runs[0].font.size = Pt(8.5)
        p0.runs[0].font.bold = True
        p1 = c1.paragraphs[0]
        p1.runs[0].font.name = "Arial"
        p1.runs[0].font.size = Pt(8.5)
        p1.runs[0].font.color.rgb = RGBColor(51, 65, 85)

    # Header row for table
    h_c0 = road_table.cell(0, 0)
    h_c1 = road_table.cell(0, 1)
    h_c0.text = "Incubation Phase"
    h_c1.text = "Deliverables & Key Milestones"
    set_cell_background(h_c0, "#0F172A")
    set_cell_background(h_c1, "#0F172A")
    for hc in [h_c0, h_c1]:
        p = hc.paragraphs[0]
        p.runs[0].font.name = "Arial"
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    doc.add_paragraph()
    add_body_p(
        "Institutional Anchoring: Hosted at the Nasarawa State University Keffi Mine-Tech UniPod, in strategic partnership with the United Nations Development Programme (UNDP) timbuktoo Africa Innovation Foundation.",
        bold_prefix="Institutional Framework: "
    )

    doc.save(output_path)
    print(f"Successfully created Concept Note DOCX at: {output_path}")

if __name__ == "__main__":
    out_dir = r"C:\Users\USER\.gemini\antigravity\scratch\regolith-sense"
    out_file = os.path.join(out_dir, "RegolithSense_Solution_Concept_Note.docx")
    create_concept_note_docx(out_file)
